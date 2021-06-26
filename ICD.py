import pandas as pd
import openpyxl as px
from openpyxl import load_workbook


wd = px.load_workbook(filename = "Book1.xlsx") #Excelファイルを読み込む
ws = wd["Sheet1"]



chart = px.chart.LineChart()
data = px.chart.Reference(ws,min_col = 2,max_col =10, min_row=4, max_row=len("群")+4)
categories = px.chart.Reference(ws, min_col=1, max_col=1, min_row=4, max_row=len("群")+4)
chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
ws.add_chart(chart, "G2")
